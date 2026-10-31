import io

import pytest
from openpyxl import Workbook, load_workbook

from tests.conftest import API as API_SEG
from tests.test_maestros import crear_empresa, usuario_con_rol

IMP = "/api/v1/importacion"
MAE = "/api/v1/maestros"
CNT = "/api/v1/contabilidad"


@pytest.fixture
def h(client, admin):
    for modulo in ("maestros", "contabilidad"):
        client.patch(f"{API_SEG}/modulos/{modulo}", headers=admin, json={"activo": True})
    return admin


def excel(encabezados, *filas) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Datos"
    hoja.append(encabezados)
    for f in filas:
        hoja.append(list(f))
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def enviar(client, h, tipo, contenido, accion="validar", empresa=None):
    datos = {"empresa_id": empresa["id"]} if empresa else {}
    return client.post(
        f"{IMP}/{tipo}/{accion}",
        headers=h,
        files={"archivo": ("carga.xlsx", contenido)},
        data=datos,
    )


def test_plantilla_con_columnas_e_instrucciones(client, h):
    r = client.get(f"{IMP}/articulos/plantilla", headers=h)
    assert r.status_code == 200
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames == ["Datos", "Instrucciones"]
    encabezados = [c.value for c in libro["Datos"][1]]
    assert encabezados[:5] == ["codigo", "descripcion", "tipo", "categoria", "unidad"]


def test_vista_previa_no_guarda_nada(client, h):
    archivo = excel(["codigo", "nombre", "codigo_padre"], ["EPP", "Protección personal", "SUM"])
    r = enviar(client, h, "categorias", archivo)
    assert r.json()["total"] == 1 and r.json()["errores"] == 0 and not r.json()["cargado"]
    codigos = [c["codigo"] for c in client.get(f"{MAE}/categorias", headers=h).json()]
    assert "EPP" not in codigos


def test_carga_todo_o_nada(client, h):
    archivo = excel(
        ["codigo", "nombre", "codigo_padre"],
        ["EPP", "Protección personal", "SUM"],
        ["EPP-CAS", "Cascos", "EPP"],  # depende de la fila anterior
        ["SUM", "Repetida", None],  # ya existe
    )
    r = enviar(client, h, "categorias", archivo, "cargar")
    assert r.json()["errores"] == 1 and not r.json()["cargado"]
    assert r.json()["filas"][0] == {
        "fila": 4,
        "clave": "SUM",
        "error": "Ya existe la categoría SUM",
    }
    codigos = [c["codigo"] for c in client.get(f"{MAE}/categorias", headers=h).json()]
    assert "EPP" not in codigos and "EPP-CAS" not in codigos
    # corregido el archivo, se carga todo
    corregido = excel(
        ["codigo", "nombre", "codigo_padre"],
        ["EPP", "Protección personal", "SUM"],
        ["EPP-CAS", "Cascos", "EPP"],
    )
    r = enviar(client, h, "categorias", corregido, "cargar")
    assert r.json()["cargado"] and r.json()["validas"] == 2
    cats = {c["codigo"]: c for c in client.get(f"{MAE}/categorias", headers=h).json()}
    assert cats["EPP-CAS"]["nivel"] == 3


def test_articulos_usan_las_mismas_reglas(client, h):
    crear_empresa(client, h)
    archivo = excel(
        [
            "codigo",
            "descripcion",
            "tipo",
            "categoria",
            "unidad",
            "precio_referencial",
            "moneda_precio",
        ],
        [None, "Papel bond A4", "Bien no inventariable", "SUM-OFI", "PAQ", 15.5, "PEN"],
        [None, "Lapicero azul", "Bien no inventariable", "SUM-OFI", "UND", None, None],
        [None, "Consultoría", "Servicio", "SUM-OFI", "SRV", None, None],  # tipo distinto al perfil
        [None, "Cinta", "Bien no inventariable", "NO-EXISTE", "UND", None, None],
        [None, "Tóner", "Bien no inventariable", "SUM-OFI", "UND", "caro", None],
    )
    r = enviar(client, h, "articulos", archivo, "cargar").json()
    assert (r["total"], r["errores"], r["cargado"]) == (5, 3, False)
    errores = {f["fila"]: f["error"] for f in r["filas"] if f["error"]}
    assert "Consumo directo" in errores[4]
    assert "NO-EXISTE" in errores[5]
    assert "no es un número" in errores[6]
    validas = enviar(
        client,
        h,
        "articulos",
        excel(
            ["descripcion", "tipo", "categoria", "unidad", "precio_referencial"],
            ["Papel bond A4", "bien no inventariable", "SUM-OFI", "PAQ", "15,50"],
            ["Lapicero azul", "bien_no_inventariable", "SUM-OFI", "UND", None],
        ),
        "cargar",
    ).json()
    assert validas["cargado"]
    arts = client.get(f"{MAE}/articulos", headers=h).json()
    assert [a["codigo"] for a in arts] == ["SUM-OFI-0001", "SUM-OFI-0002"]
    assert arts[0]["precio_referencial"] == "15.5000" and arts[0]["moneda_precio"] == "PEN"


def test_proveedores_con_contacto_y_dni_sin_ceros(client, h):
    archivo = excel(
        [
            "tipo_documento",
            "numero_documento",
            "razon_social",
            "condicion_pago",
            "contacto_nombre",
            "contacto_correo",
        ],
        ["RUC", 20100047218, "Banco de Crédito", "30D", "Ana Ruiz", "ana@bcp.pe"],
        ["DNI", 4567891, "Juan Pérez", None, None, None],  # Excel quitó el cero inicial
        ["RUC", "20100047219", "RUC malo", None, None, None],
    )
    r = enviar(client, h, "proveedores", archivo).json()
    assert r["errores"] == 1 and "dígito verificador" in r["filas"][0]["error"]
    r = enviar(
        client,
        h,
        "proveedores",
        excel(
            [
                "tipo_documento",
                "numero_documento",
                "razon_social",
                "condicion_pago",
                "contacto_nombre",
            ],
            ["RUC", 20100047218, "Banco de Crédito", "30D", "Ana Ruiz"],
            ["DNI", 4567891, "Juan Pérez", None, None],
        ),
        "cargar",
    ).json()
    assert r["cargado"]
    provs = {p["numero_documento"]: p for p in client.get(f"{MAE}/proveedores", headers=h).json()}
    assert "04567891" in provs and provs["20100047218"]["condicion_pago"]["codigo"] == "30D"
    ficha = client.get(f"{MAE}/proveedores/{provs['20100047218']['id']}", headers=h).json()
    assert ficha["contactos"][0]["nombre"] == "Ana Ruiz" and ficha["contactos"][0]["principal"]


def test_plan_de_cuentas_por_empresa(client, h):
    e = crear_empresa(client, h).json()
    archivo = excel(
        ["codigo", "nombre"], ["6321", "Asesoría y consultoría"], ["63211", "Asesoría legal"]
    )
    assert (
        enviar(client, h, "plan_cuentas", archivo, "cargar").json()["codigo"] == "empresa_requerida"
    )
    r = enviar(client, h, "plan_cuentas", archivo, "cargar", empresa=e).json()
    assert r["cargado"]
    cuentas = {
        c["codigo"]: c for c in client.get(f"{CNT}/empresas/{e['id']}/cuentas", headers=h).json()
    }
    assert (
        cuentas["63211"]["padre_id"] == cuentas["6321"]["id"] and not cuentas["6321"]["imputable"]
    )


def test_permisos_y_acceso_a_la_empresa(client, h):
    e = crear_empresa(client, h).json()
    archivo = excel(["codigo", "nombre"], ["6321", "Asesoría"])
    consulta = usuario_con_rol(client, h, "ERP_MAE_CONSULTA")
    assert (
        enviar(client, consulta, "categorias", excel(["codigo", "nombre"], ["X", "Y"])).status_code
        == 403
    )
    contador = usuario_con_rol(client, h, "ERP_CNT_ADMINISTRADOR")  # sin empresas asignadas
    r = enviar(client, contador, "plan_cuentas", archivo, empresa=e)
    assert r.json()["codigo"] == "sin_acceso_empresa"


def test_archivo_invalido_o_sin_columnas(client, h):
    assert enviar(client, h, "categorias", b"no es excel").json()["codigo"] == "archivo_invalido"
    assert (
        enviar(client, h, "categorias", excel(["nombre"], ["Sin código"])).json()["codigo"]
        == "columnas_faltantes"
    )
    assert (
        enviar(client, h, "categorias", excel(["codigo", "nombre"])).json()["codigo"]
        == "archivo_vacio"
    )
    assert client.get(f"{IMP}/inventado/plantilla", headers=h).status_code == 404


def test_carga_queda_auditada(client, h):
    enviar(client, h, "categorias", excel(["codigo", "nombre"], ["EPP", "Protección"]), "cargar")
    eventos = client.get(f"{API_SEG}/auditoria?entidad=importacion", headers=h).json()
    assert eventos[0]["detalle"] == {"tipo": "categorias", "filas": 1}
