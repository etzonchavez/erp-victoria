"""Carga de maestros desde Excel: plantilla, vista previa y carga todo o nada.

Cada fila se procesa con las mismas reglas que la pantalla (los servicios de cada módulo),
dentro de un savepoint para reunir todos los errores. Si alguna fila falla no se carga
nada; si todas están bien, se confirma la carga completa en una sola transacción.
"""

import io
import unicodedata
import uuid
from collections.abc import Callable
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import ErrorERP
from app.modules.contabilidad import service as contabilidad
from app.modules.contabilidad.schemas import CuentaCrear
from app.modules.maestros import datos_base
from app.modules.maestros import service as maestros
from app.modules.maestros.models import Categoria, CondicionPago
from app.modules.maestros.schemas import (
    ArticuloCrear,
    CategoriaCrear,
    ContactoCrear,
    ProveedorCrear,
)
from app.modules.seguridad import service as seguridad
from app.modules.seguridad.models import Usuario
from app.modules.seguridad.service import auditar

MAX_FILAS = 5000
MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS_RESPUESTA = 500


@dataclass(frozen=True)
class Columna:
    nombre: str
    descripcion: str
    obligatoria: bool = False
    ejemplo: str = ""


@dataclass(frozen=True)
class TipoCarga:
    codigo: str
    nombre: str
    permiso: str
    columnas: tuple[Columna, ...]
    por_empresa: bool = False


TIPOS: dict[str, TipoCarga] = {
    t.codigo: t
    for t in (
        TipoCarga(
            "categorias",
            "Categorías de artículo",
            "maestros.categoria.editar",
            (
                Columna("codigo", "Código único (mayúsculas, números, - o _)", True, "SUM-EPP"),
                Columna("nombre", "Nombre de la categoría", True, "Equipos de protección"),
                Columna(
                    "codigo_padre",
                    "Código de la categoría padre (vacío = primer nivel). Puede estar más "
                    "arriba en el mismo archivo",
                    False,
                    "SUM",
                ),
            ),
        ),
        TipoCarga(
            "articulos",
            "Artículos",
            "maestros.articulo.crear",
            (
                Columna("codigo", "Vacío = correlativo automático de la categoría", False, ""),
                Columna("descripcion", "Descripción", True, "Casco de seguridad blanco"),
                Columna(
                    "tipo",
                    "Bien inventariable, Bien no inventariable, Servicio o Activo fijo",
                    True,
                    "Bien inventariable",
                ),
                Columna("categoria", "Código de la categoría", True, "SUM-EPP"),
                Columna("unidad", "Código de la unidad de medida", True, "UND"),
                Columna("impuesto", "Código del impuesto (vacío = IGV18)", False, "IGV18"),
                Columna("marca", "Marca", False, "3M"),
                Columna("modelo", "Modelo", False, "H-701"),
                Columna("especificacion", "Especificación técnica", False, "Norma ANSI Z89.1"),
                Columna("precio_referencial", "Precio estimado (número)", False, "45.90"),
                Columna("moneda_precio", "Moneda del precio (PEN, USD…)", False, "PEN"),
            ),
        ),
        TipoCarga(
            "proveedores",
            "Proveedores",
            "maestros.proveedor.crear",
            (
                Columna("tipo_documento", "RUC, DNI o EXT (exterior)", True, "RUC"),
                Columna("numero_documento", "Número de documento", True, "20100047218"),
                Columna("pais", "Solo proveedores del exterior: código de 2 letras", False, ""),
                Columna("razon_social", "Razón social", True, "Banco de Crédito del Perú"),
                Columna("nombre_comercial", "Nombre comercial", False, "BCP"),
                Columna("direccion", "Dirección", False, "Av. Centenario 156, La Molina"),
                Columna("correo", "Correo", False, "ventas@proveedor.pe"),
                Columna("telefono", "Teléfono", False, "01 311 9898"),
                Columna(
                    "moneda", "Moneda habitual (vacío = PEN o USD si es del exterior)", False, "PEN"
                ),
                Columna("condicion_pago", "Código de la condición (vacío = CONTADO)", False, "30D"),
                Columna("contacto_nombre", "Contacto principal: nombre", False, "Ana Ruiz"),
                Columna("contacto_cargo", "Contacto principal: cargo", False, "Ejecutiva"),
                Columna("contacto_correo", "Contacto principal: correo", False, "ana@proveedor.pe"),
                Columna("contacto_telefono", "Contacto principal: teléfono", False, "987654321"),
            ),
        ),
        TipoCarga(
            "plan_cuentas",
            "Plan de cuentas",
            "contabilidad.cuenta.editar",
            (
                Columna(
                    "codigo", "Código numérico; el padre se detecta por el prefijo", True, "63211"
                ),
                Columna("nombre", "Nombre de la cuenta", True, "Asesoría legal"),
                Columna(
                    "naturaleza",
                    "Solo cuentas de primer nivel: activo, pasivo, patrimonio, gasto, ingreso…",
                    False,
                    "",
                ),
            ),
            por_empresa=True,
        ),
    )
}


# ---------------------------------------------------------------- plantilla


def plantilla(tipo: str) -> bytes:
    t = _tipo(tipo)
    libro = Workbook()
    datos = libro.active
    datos.title = "Datos"
    negrita, relleno = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E2A4A")
    for i, col in enumerate(t.columnas, start=1):
        celda = datos.cell(row=1, column=i, value=col.nombre)
        celda.font, celda.fill = negrita, relleno
        datos.column_dimensions[get_column_letter(i)].width = max(14, len(col.nombre) + 4)
        # Los códigos y documentos se guardan como texto para no perder ceros a la izquierda.
        for fila in range(2, 202):
            datos.cell(row=fila, column=i).number_format = "@"
    datos.freeze_panes = "A2"
    ayuda = libro.create_sheet("Instrucciones")
    ayuda.append([f"Carga de {t.nombre.lower()} — Victoria ERP"])
    ayuda["A1"].font = Font(bold=True, size=13)
    ayuda.append(
        ["Llena la hoja Datos desde la fila 2. Si una fila tiene un error no se carga nada."]
    )
    ayuda.append([])
    ayuda.append(["Columna", "Obligatoria", "Descripción", "Ejemplo"])
    for c in ayuda[4]:
        c.font, c.fill = negrita, relleno
    for col in t.columnas:
        ayuda.append([col.nombre, "Sí" if col.obligatoria else "No", col.descripcion, col.ejemplo])
    for letra, ancho in zip("ABCD", (22, 12, 80, 30), strict=True):
        ayuda.column_dimensions[letra].width = ancho
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------- lectura


def _normalizar(texto) -> str:
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return texto.strip().lower().replace(" ", "_").replace("°", "")


def _celda(valor) -> str | None:
    """Texto limpio de una celda: los números enteros pierden el '.0' de Excel."""
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    return texto or None


def leer(tipo: str, contenido: bytes) -> list[tuple[int, dict]]:
    t = _tipo(tipo)
    if len(contenido) > MAX_BYTES:
        raise ErrorERP(413, "archivo_grande", "El archivo no puede superar 5 MB")
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise ErrorERP(
            422, "archivo_invalido", "El archivo no es un Excel válido (.xlsx)"
        ) from None
    hoja = libro["Datos"] if "Datos" in libro.sheetnames else libro.worksheets[0]
    filas = hoja.iter_rows(values_only=True)
    encabezados = [_normalizar(v) for v in next(filas, ())]
    esperadas = {c.nombre for c in t.columnas}
    faltan = [c.nombre for c in t.columnas if c.obligatoria and c.nombre not in encabezados]
    if faltan:
        raise ErrorERP(
            422,
            "columnas_faltantes",
            f"Faltan columnas obligatorias: {', '.join(faltan)}. Descarga la plantilla.",
        )
    resultado = []
    for numero, fila in enumerate(filas, start=2):
        datos = {
            encabezados[i]: _celda(v)
            for i, v in enumerate(fila)
            if i < len(encabezados) and encabezados[i] in esperadas
        }
        if not any(datos.values()):
            continue
        resultado.append((numero, datos))
        if len(resultado) > MAX_FILAS:
            raise ErrorERP(422, "demasiadas_filas", f"El archivo tiene más de {MAX_FILAS} filas")
    if not resultado:
        raise ErrorERP(422, "archivo_vacio", "La hoja Datos no tiene filas para cargar")
    return resultado


# ---------------------------------------------------------------- procesamiento


def _tipo(tipo: str) -> TipoCarga:
    t = TIPOS.get(tipo)
    if not t:
        raise ErrorERP(404, "tipo_invalido", f"Tipo de carga no válido: {', '.join(TIPOS)}")
    return t


def _validar(esquema: type[BaseModel], datos: dict) -> dict:
    try:
        return esquema.model_validate(datos).model_dump()
    except ValidationError as e:
        err = e.errors()[0]
        campo = ".".join(str(p) for p in err["loc"]) or "fila"
        raise ErrorERP(422, "validacion", f"{campo}: {err['msg']}") from None


def _categoria_por_codigo(db: Session, codigo: str | None) -> uuid.UUID | None:
    if not codigo:
        return None
    cat = db.scalar(select(Categoria).where(Categoria.codigo == codigo.strip().upper()))
    if not cat:
        raise ErrorERP(422, "categoria_inexistente", f"No existe la categoría {codigo}")
    return cat.id


TIPOS_POR_NOMBRE = {_normalizar(n): c for c, n in datos_base.TIPOS_ARTICULO.items()} | {
    c: c for c in datos_base.TIPOS_ARTICULO
}


def _fila_categoria(db, actor, d, _empresa):
    datos = _validar(
        CategoriaCrear,
        {"codigo": d.get("codigo"), "nombre": d.get("nombre"), "padre_id": None},
    )
    padre = _categoria_por_codigo(db, d.get("codigo_padre"))
    maestros.crear_categoria(db, actor, datos["codigo"], datos["nombre"], padre, commit=False)
    return datos["codigo"].upper()


def _fila_articulo(db, actor, d, _empresa):
    tipo = TIPOS_POR_NOMBRE.get(_normalizar(d.get("tipo")))
    if not tipo:
        raise ErrorERP(
            422, "tipo_invalido", f"Tipo no válido: {', '.join(datos_base.TIPOS_ARTICULO.values())}"
        )
    precio = d.get("precio_referencial")
    datos = _validar(
        ArticuloCrear,
        {
            **{
                k: d.get(k)
                for k in (
                    "codigo",
                    "descripcion",
                    "unidad",
                    "impuesto",
                    "marca",
                    "modelo",
                    "especificacion",
                    "moneda_precio",
                )
            },
            "tipo": tipo,
            "categoria_id": _categoria_por_codigo(db, d.get("categoria")) or _falta("categoria"),
            "precio_referencial": _numero(precio, "precio_referencial"),
        },
    )
    if datos.get("precio_referencial") is not None and not datos.get("moneda_precio"):
        datos["moneda_precio"] = "PEN"
    return maestros.crear_articulo(db, actor, datos, commit=False).codigo


def _fila_proveedor(db, actor, d, _empresa):
    condicion = None
    if d.get("condicion_pago"):
        condicion = db.scalar(
            select(CondicionPago).where(CondicionPago.codigo == d["condicion_pago"].upper())
        )
        if not condicion:
            raise ErrorERP(
                422,
                "condicion_inexistente",
                f"No existe la condición de pago {d['condicion_pago']}",
            )
    numero = d.get("numero_documento") or ""
    if (d.get("tipo_documento") or "").upper() == "DNI" and numero.isdigit():
        numero = numero.zfill(8)  # Excel quita los ceros iniciales de los DNI
    datos = _validar(
        ProveedorCrear,
        {
            **{
                k: d.get(k)
                for k in (
                    "tipo_documento",
                    "pais",
                    "razon_social",
                    "nombre_comercial",
                    "direccion",
                    "correo",
                    "telefono",
                    "moneda",
                )
            },
            "numero_documento": numero,
            "condicion_pago_id": condicion.id if condicion else None,
        },
    )
    prov = maestros.crear_proveedor(db, actor, datos, commit=False)
    if d.get("contacto_nombre"):
        contacto = _validar(
            ContactoCrear,
            {
                "nombre": d["contacto_nombre"],
                "cargo": d.get("contacto_cargo"),
                "correo": d.get("contacto_correo"),
                "telefono": d.get("contacto_telefono"),
                "principal": True,
            },
        )
        maestros.crear_contacto(db, actor, prov.id, contacto, commit=False)
    return f"{prov.tipo_documento} {prov.numero_documento}"


def _fila_cuenta(db, actor, d, empresa):
    datos = _validar(CuentaCrear, {k: d.get(k) for k in ("codigo", "nombre", "naturaleza")})
    contabilidad.crear_cuenta(
        db, actor, empresa, datos["codigo"], datos["nombre"], datos.get("naturaleza"), commit=False
    )
    return datos["codigo"]


def _numero(texto: str | None, campo: str) -> Decimal | None:
    if not texto:
        return None
    try:
        return Decimal(texto.replace(",", "."))
    except InvalidOperation:
        raise ErrorERP(422, "valor_invalido", f"{campo}: {texto} no es un número") from None


def _falta(campo: str):
    raise ErrorERP(422, "dato_requerido", f"Falta {campo}")


PROCESADORES: dict[str, Callable] = {
    "categorias": _fila_categoria,
    "articulos": _fila_articulo,
    "proveedores": _fila_proveedor,
    "plan_cuentas": _fila_cuenta,
}


def procesar(
    db: Session,
    actor: Usuario,
    tipo: str,
    contenido: bytes,
    empresa_id: uuid.UUID | None,
    confirmar: bool,
) -> dict:
    """Valida todas las filas; si `confirmar` y no hay errores, carga todo."""
    t = _tipo(tipo)
    modulo = t.permiso.split(".")[0]
    if modulo not in seguridad.modulos_activos(db):
        raise ErrorERP(403, "modulo_inactivo", f"El módulo {modulo} no está activo")
    if t.permiso not in seguridad.permisos_de(db, actor):
        raise ErrorERP(403, "sin_permiso", f"No tienes el permiso {t.permiso}")
    if t.por_empresa:
        if not empresa_id:
            raise ErrorERP(422, "empresa_requerida", "Elige la empresa")
        maestros.obtener_empresa(db, empresa_id)
        maestros.verificar_acceso_empresa(db, actor, empresa_id)
    filas = leer(tipo, contenido)
    resultado = []
    for numero, datos in filas:
        try:
            with db.begin_nested():
                clave = PROCESADORES[tipo](db, actor, datos, empresa_id)
            resultado.append({"fila": numero, "clave": clave, "error": None})
        except ErrorERP as e:
            resultado.append({"fila": numero, "clave": _clave_visible(datos), "error": e.mensaje})
    errores = sum(1 for r in resultado if r["error"])
    cargado = confirmar and errores == 0
    if cargado:
        auditar(
            db,
            actor,
            "importar",
            "importacion",
            empresa_id,
            {"tipo": tipo, "filas": len(filas)},
        )
        db.commit()
    else:
        db.rollback()
    con_error = [r for r in resultado if r["error"]]
    sin_error = [r for r in resultado if not r["error"]]
    return {
        "tipo": tipo,
        "total": len(resultado),
        "validas": len(resultado) - errores,
        "errores": errores,
        "cargado": cargado,
        # primero las filas con error, para que se vean de inmediato
        "filas": (con_error + sin_error)[:MAX_FILAS_RESPUESTA],
    }


def _clave_visible(datos: dict) -> str:
    for campo in ("codigo", "numero_documento", "descripcion", "razon_social", "nombre"):
        if datos.get(campo):
            return datos[campo]
    return "—"
